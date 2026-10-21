#!/usr/bin/env python3
"""build-xlsx.py - 把 parsed-v3.json 写成采购表 Excel（openpyxl，只读数据、不联网）。"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "parsed-v3.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else "out/淘宝购物车采购表.xlsx"
d = json.load(open(SRC, encoding="utf-8"))
items = d["items"]

wb = Workbook()
ws = wb.active
ws.title = "采购表"

HEAD = ["序号", "店铺", "商品名称", "规格/颜色/容量", "单价(¥)", "原价(¥)", "数量", "小计(¥)", "备注"]
WIDTH = [6, 22, 58, 30, 11, 11, 6, 12, 18]
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F4E79")

for c, (h, w) in enumerate(zip(HEAD, WIDTH), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 22

r = 2
for n, it in enumerate(items, 1):
    price = it.get("price") or 0
    qty = it.get("qty") or 0
    ws.cell(row=r, column=1, value=n).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=2, value=it.get("shop") or "")
    ws.cell(row=r, column=3, value=it.get("title") or "")
    ws.cell(row=r, column=4, value=it.get("spec") or "")
    c5 = ws.cell(row=r, column=5, value=price); c5.number_format = '#,##0.00'
    if it.get("orig_price"):
        c6 = ws.cell(row=r, column=6, value=it["orig_price"]); c6.number_format = '#,##0.00'
    ws.cell(row=r, column=7, value=qty).alignment = Alignment(horizontal="center")
    c8 = ws.cell(row=r, column=8, value=f"=E{r}*G{r}"); c8.number_format = '#,##0.00'
    marks = []
    if not it.get("orig_price"):
        marks.append("原价缺失")
    if not (it.get("shop") or "").strip():
        marks.append("店铺未识别")
    ws.cell(row=r, column=9, value="；".join(marks))
    for c in range(1, 10):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = Alignment(
            horizontal=ws.cell(row=r, column=c).alignment.horizontal or "general",
            vertical="center", wrap_text=(c in (3, 4)))
    r += 1

last = r - 1
ws.cell(row=r, column=1, value="合计").font = Font(bold=True)
ws.cell(row=r, column=7, value=f"=SUM(G2:G{last})").font = Font(bold=True)
ws.cell(row=r, column=8, value=f"=SUM(H2:H{last})").font = Font(bold=True)
ws.cell(row=r, column=8).number_format = '#,##0.00'
for c in range(1, 10):
    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
    ws.cell(row=r, column=c).border = border
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:I{last}"

note = wb.create_sheet("说明")
rows = [
    ("数据来源", "Edge 浏览器已登录会话中的 https://cart.taobao.com/cart.htm（账号 tb455443353821），2026-10-21 14:40–15:30 CST"),
    ("采集方式", "纯 GUI 驱动：Ctrl+A/Ctrl+C 取页面渲染文本（滚动懒加载逐轮累积）+ 12 张 50% 缩放截图做视觉交叉核对；未用接口抓包、未用 DevTools 协议"),
    ("行数", f"{len(items)} 行（页面文本中的购物车条目；同一商品不同 SKU 视为多行，未去重）"),
    ("金额", f"本表单价×数量合计 ¥{d['amount_sum']:,.2f}；数量合计 {d['qty_sum']} 件"),
    ("对账差异（重要）", "页面全选后显示「领券结算(50)」合计 ¥264,972.18，与购物车角标「138」、页面文本 62 个条目均不一致。可能原因：角标统计件数而非条目数；部分条目不可结算；页面存在隐藏/失效条目。本表不做推测性修正，差异如实登记"),
    ("单价可信度", "抽查 6/6 行与截图一致（¥37.3/¥37190/¥8080/¥3898.3/¥5.51/¥30.11）"),
    ("原价可信度", "中。部分行原价被「距加入降」金额干扰（如 Lexar 原价 ¥6999 已修正，其余标记「原价缺失」）"),
    ("数量可信度", "中高。抽查发现 1 处不一致（鹿小班 INMP441 文本层为 1，截图显示 2）"),
    ("为什么没有状态列", "页面文本中「重新选择规格」「宝贝已不能购买」等字样在每一行模板里都存在（CSS 隐藏），无法用文本区分真实状态；可靠判定需要逐行视觉识别。截图中确认不可购买的条目已在下方列出，供人工复核"),
    ("视觉确认不可结算的条目", "宏碁掠夺者 GM9000【8TB】、iBook Pocket5、GPD Pocket3（多行）、ZWIDE N8 之外的若干「宝贝已不能购买」行：悦享 ¥159.9、绿联 2025 新款、NORGE 5C 800-1600W ¥649、Apple Mac Studio 16 核 ¥7899、ITECH 电池测试仪 ¥8200、FEVM FA-X1、阿迈奇 F5A 等（以截图 s01–s12 为准）"),
    ("列说明", "单价=页面「平台加补后」价；小计=单价×数量；规格列来自页面 SKU 文本，缺失者为空"),
    ("生成方式", "missions/taobao-cart-2026-10-21/build-xlsx.py（openpyxl），输入 parsed-v3.json"),
]
note.column_dimensions["A"].width = 26
note.column_dimensions["B"].width = 110
for i, (k, v) in enumerate(rows, 1):
    a = note.cell(row=i, column=1, value=k); a.font = Font(bold=True)
    a.alignment = Alignment(vertical="top")
    b = note.cell(row=i, column=2, value=v)
    b.alignment = Alignment(wrap_text=True, vertical="top")

import os
os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
wb.save(OUT)
print(json.dumps({"ok": True, "out": OUT, "rows": len(items), "amount": d["amount_sum"]}, ensure_ascii=False))
